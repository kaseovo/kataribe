"""Fills the AI Builder Cup submission template with the Kataribe pitch.

    uv run --no-project --with python-pptx --with "qrcode[pil]" python docs/demo/deck/build.py

The template is prescribed: submissions that do not follow it may be
disqualified. So every slide and heading it ships with stays where it is, and
only its "download this template" note is removed. Everything else is drawn on
top, in the template's own font.

Reads template.pptx and assets/ beside this file; writes kataribe-pitch-deck.pptx
one level up. Edit the wording here, not in the output.
"""

import copy
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
ICONS = HERE / "assets" / "icons"
SHOTS = HERE / "assets" / "shots"
OUT = HERE.parent / "kataribe-pitch-deck.pptx"

# Facts still to confirm or fill in before submitting.
TEAM_NAME = "Kataribe"
TEAM_LEADER = "Tanguy Pemeja"
PROBLEM_STATEMENT = "Sustainability & Social Impact"
REPO_URL = "https://github.com/kaseovo/kataribe"
VIDEO_URL = "https://youtu.be/wg3SdraT2Oo"  # public on YouTube, 3:48
PRODUCT_URL = "https://kataribe.kaseovo.com"

INK = "202124"
MUTED = "5F6368"
CARD = "F1F3F4"
LINE = "DADCE0"
DARK = "1F1F1F"
ACCENT = "D98F4E"  # the app's own accent
ACCENT_TEXT = "A0561B"  # the accent, dark enough for small text on white
ACCENT_TINT = "FBEEE2"
ON_DARK = "F0EBE5"
ON_DARK_MUTED = "B5ABA1"
APP_BG = "12100E"
APP_PANEL = "26211D"
WHITE = "FFFFFF"

FONTS = {
    "regular": "Google Sans Flex",
    "medium": "Google Sans Flex Medium",
    "semibold": "Google Sans Flex SemiBold",
}
# On every Mac; elsewhere it falls back to the system's own Japanese sans-serif.
JAPANESE = "Hiragino Sans"

# The content area under the template's heading, in inches.
LEFT, RIGHT, TOP = 0.39, 9.61, 1.45
WIDTH = RIGHT - LEFT


def emu(inches: float) -> Emu:
    return Emu(int(round(inches * 914400)))


# --- text ---------------------------------------------------------------------


def R(text, size=12, weight="regular", color=INK, link=None):
    return {"text": text, "size": size, "weight": weight, "color": color, "link": link}


def P(*runs, align="l", after=0, before=0, bullet=False, line=None):
    return {"runs": runs, "align": align, "after": after, "before": before, "bullet": bullet, "line": line}


def _fill(parent, color):
    solid = etree.SubElement(parent, qn("a:solidFill"))
    etree.SubElement(solid, qn("a:srgbClr")).set("val", color)


UNIT = re.compile(r"(\d) (?=(?:s|ms|min|minutes|kHz|MB|MiB|GB|million)\b)")


def _write_run(paragraph, spec):
    run = paragraph.add_run()
    # A figure never ends a line without its unit.
    run.text = UNIT.sub("\\1\u00a0", spec["text"])
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "en-GB")
    rPr.set("sz", str(int(spec["size"] * 100)))
    rPr.set("dirty", "0")
    _fill(rPr, spec["color"])
    face = FONTS[spec["weight"]]
    for tag, typeface in (("a:latin", face), ("a:ea", JAPANESE), ("a:cs", face), ("a:sym", face)):
        etree.SubElement(rPr, qn(tag)).set("typeface", typeface)
    if spec["link"]:
        run.hyperlink.address = spec["link"]


def _write_paragraph(paragraph, spec, bullet_color=ACCENT):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("algn", spec["align"])
    if spec["line"]:
        etree.SubElement(etree.SubElement(pPr, qn("a:lnSpc")), qn("a:spcPct")).set(
            "val", str(int(spec["line"] * 100000))
        )
    etree.SubElement(etree.SubElement(pPr, qn("a:spcBef")), qn("a:spcPts")).set("val", str(int(spec["before"] * 100)))
    etree.SubElement(etree.SubElement(pPr, qn("a:spcAft")), qn("a:spcPts")).set("val", str(int(spec["after"] * 100)))
    if spec["bullet"]:
        size = spec["runs"][0]["size"]
        pPr.set("marL", str(int(Pt(size * 1.1))))
        pPr.set("indent", str(-int(Pt(size * 1.1))))
        _fill(etree.SubElement(pPr, qn("a:buClr")), bullet_color)
        etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
    else:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))
    for run in spec["runs"]:
        _write_run(paragraph, run)


def fill_frame(frame, paragraphs, anchor="t", inset=0.0, bullet_color=ACCENT):
    frame.word_wrap = True
    body = frame._txBody.bodyPr
    for side in ("lIns", "tIns", "rIns", "bIns"):
        body.set(side, str(int(emu(inset))))
    body.set("anchor", anchor)
    for child in list(body):
        body.remove(child)
    etree.SubElement(body, qn("a:noAutofit"))
    first = True
    for spec in paragraphs:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        _write_paragraph(paragraph, spec, bullet_color)


def text(slide, x, y, w, h, paragraphs, anchor="t", name=None, bullet_color=ACCENT):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    if name:
        box.name = name
    fill_frame(box.text_frame, paragraphs, anchor, bullet_color=bullet_color)
    return box


# --- shapes -------------------------------------------------------------------


def _plain(shape):
    """Drops the theme style python-pptx attaches, so only explicit looks apply."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def rect(slide, x, y, w, h, fill, radius=0.1, line=None, name=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, emu(x), emu(y), emu(w), emu(h))
    _plain(s)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    if line:
        s.line.color.rgb = _rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    if name:
        s.name = name
    return s


def circle(slide, x, y, d, fill, name=None):
    return rect(slide, x, y, d, d, fill, name=name, shape=MSO_SHAPE.OVAL)


def _rgb(hex_color):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_color)


def label_in(shape, paragraphs, anchor="ctr", inset=0.04):
    fill_frame(shape.text_frame, paragraphs, anchor, inset)
    return shape


def icon(slide, name, x, y, size, color=ACCENT):
    return slide.shapes.add_picture(str(ICONS / f"{name}-{color}.png"), emu(x), emu(y), emu(size), emu(size))


def badge(slide, name, x, y, d=0.46):
    """The deck's motif: an amber icon in a dark circle."""
    circle(slide, x, y, d, DARK)
    pad = d * 0.22
    icon(slide, name, x + pad, y + pad, d - 2 * pad)


def number(slide, n, x, y, d=0.3, fill=DARK, color=ACCENT):
    c = circle(slide, x, y, d, fill)
    return label_in(c, [P(R(str(n), 10, "semibold", color), align="ctr")], inset=0)


def arrow(slide, x1, y1, x2, y2, color="9AA0A6", width=1.5, both=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    c.line.color.rgb = _rgb(color)
    c.line.width = Pt(width)
    _plain(c)
    ln = c.line._get_or_add_ln()
    if both:
        etree.SubElement(ln, qn("a:headEnd")).attrib.update({"type": "triangle", "w": "med", "len": "med"})
    etree.SubElement(ln, qn("a:tailEnd")).attrib.update({"type": "triangle", "w": "med", "len": "med"})
    return c


def picture(slide, path, x, y, w, h, radius=0.0):
    pic = slide.shapes.add_picture(str(path), emu(x), emu(y), emu(w), emu(h))
    if radius:
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        etree.SubElement(av, qn("a:gd")).attrib.update({"name": "adj", "fmla": f"val {int(radius / min(w, h) * 100000)}"})
    return pic


def table(slide, x, y, widths, rows, header_fill=DARK, size=10.5, row_h=0.36):
    """Rows of strings; the first is the header. A trailing '*' bolds a row."""
    frame = slide.shapes.add_table(len(rows), len(widths), emu(x), emu(y), emu(sum(widths)), emu(row_h * len(rows)))
    tbl = frame.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for i, w in enumerate(widths):
        tbl.columns[i].width = emu(w)
    for r, row in enumerate(rows):
        tbl.rows[r].height = emu(row_h)
        strong = r == 0 or row[-1] == "*"
        cells = row[:-1] if row[-1] == "*" else row
        for c, value in enumerate(cells):
            cell = tbl.cell(r, c)
            color = WHITE if r == 0 else INK
            weight = "semibold" if strong else "regular"
            fill_frame(
                cell.text_frame,
                [P(R(value, size, weight, color), align="r" if c == len(cells) - 1 and c > 0 else "l")],
                anchor="ctr",
            )
            # Cells take their padding from the margins below, not the text body.
            for side in ("lIns", "tIns", "rIns", "bIns"):
                cell.text_frame._txBody.bodyPr.attrib.pop(side, None)
            cell.margin_left = cell.margin_right = emu(0.1)
            cell.margin_top = cell.margin_bottom = emu(0.03)
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.set("anchor", "ctr")
            for child in list(tcPr):
                tcPr.remove(child)
            ln = etree.SubElement(tcPr, qn("a:lnB"))
            ln.set("w", "6350")
            _fill(ln, LINE if r else header_fill)
            _fill(tcPr, header_fill if r == 0 else WHITE)
    return frame


# --- template -----------------------------------------------------------------


def heading_box(slide):
    """The template's own heading text box, at 0.73" from the top."""
    for shape in slide.shapes:
        if shape.has_text_frame and abs(shape.top - emu(0.73)) < emu(0.05):
            return shape
    raise LookupError("no heading box on this slide")


def remove_slide(prs, index):
    ids = prs.slides._sldIdLst
    entry = ids[index]
    prs.part.drop_rel(entry.rId)
    ids.remove(entry)


# --- slides -------------------------------------------------------------------


def team_details(slide):
    box = next(s for s in slide.shapes if s.has_text_frame and "Team name" in s.text_frame.text)
    values = {"Team name": TEAM_NAME, "Team leader name": TEAM_LEADER, "Problem Statement": PROBLEM_STATEMENT}
    for paragraph in box.text_frame.paragraphs:
        for key, value in values.items():
            if paragraph.runs and paragraph.runs[0].text.startswith(key):
                _write_run(paragraph, R(value, 18, "semibold", INK))
                # The value belongs before the paragraph's end mark.
                end = paragraph._p.find(qn("a:endParaRPr"))
                if end is not None:
                    paragraph._p.remove(end)
                    paragraph._p.append(end)


def brief(slide):
    text(
        slide, LEFT, TOP, 5.6, 0.55,
        [P(R("Kataribe ", 28, "semibold"), R("語り部", 24, "semibold", ACCENT_TEXT))],
        name="Name",
    )
    text(
        slide, LEFT, 2.02, 5.5, 1.05,
        [
            P(R("Every grandparent has stories nobody has asked about yet.", 14, "semibold"), after=4),
            P(R("Kataribe talks with an elderly person for a few minutes a day, in Japanese, and turns what "
                "they remember into a life story their family can read, and hear in their own voice.", 11.5,
                "regular", INK), line=1.05),
        ],
        name="One-liner",
    )
    rows = [
        ("THE PROBLEM", "Many Japanese seniors spend their days alone, and their families never find the "
                        "time or the courage to ask about their lives. The stories go unrecorded."),
        ("WHO IT'S FOR", "Them: a few minutes a day about their own life. Their family: the stories they "
                         "never asked for. Next, their town: a friend, a carer, a local archive."),
        ("NOT A COMPANION", "It does not stand in for anyone. A kataribe is a keeper of oral history; this "
                            "one helps the people around them know them better."),
    ]
    y = 3.12
    for label, body in rows:
        text(slide, LEFT, y, 5.5, 0.72,
             [P(R(label, 9.5, "semibold", ACCENT_TEXT), after=2), P(R(body, 11, "regular", INK), line=1.05)],
             name=label.title())
        y += 0.74

    stats = [
        ("6.72M", "people aged 65 and over live alone in Japan",
         "2020 census · Cabinet Office, Annual Report on the Ageing Society"),
        ("¥220,000", "for one professionally written self-history (自分史), made once, by hand",
         "Oya no Zasshi (親の雑誌, Kokoromi), price incl. tax"),
    ]
    y = TOP
    for big, label, source in stats:
        rect(slide, 6.25, y, RIGHT - 6.25, 1.78, DARK, radius=0.12, name=f"Stat {big}")
        text(slide, 6.5, y + 0.2, RIGHT - 6.75, 1.45, [
            P(R(big, 32, "semibold", ACCENT), after=4),
            P(R(label, 11.5, "medium", ON_DARK), after=6, line=1.05),
            P(R(source, 8.5, "regular", ON_DARK_MUTED)),
        ], name=f"Stat {big} text")
        y += 1.98


def opportunities(slide):
    box = heading_box(slide)
    paragraphs = box.text_frame._txBody.findall(qn("a:p"))
    for extra in paragraphs[1:]:
        box.text_frame._txBody.remove(extra)
    for br in paragraphs[0].findall(qn("a:br")):
        paragraphs[0].remove(br)
    box.height = emu(0.61)

    text(slide, LEFT, 1.38, WIDTH, 0.35,
         [P(R("Not a companion that stands in for anyone: a way for the people around them to know them better.",
              13, "medium", ACCENT_TEXT))],
         name="Positioning")
    cards = [
        ("How different is it from any of the other existing ideas?", [
            "Memoir AIs (Tell Mel, Storyworth, Remento) rewrite a parent's words into chapters, and "
            "the family chooses who receives them.",
            "Companion AIs (Cotomo, by Starley) keep seniors company, but no life story comes out.",
            "A human-written self-history costs ¥220,000, and is made once.",
        ]),
        ("How will it be able to solve the problem?", [
            "A few minutes a day in natural Japanese. The interviewer speaks first: nothing to install or learn.",
            "Each conversation becomes stories, set as a book by time of life. What is untold plans the next.",
            "Relatives read what was shared, hear the voice behind each quote, and send questions it asks next.",
        ]),
        ("USP of the proposed solution", [
            "Theirs to give: nothing reaches the family without a yes in their own words, and any story can "
            "be taken back.",
            "It retells, but never adds: every fact has a quote, found word for word, then heard in the recording.",
            "Built for Japanese: honorific register, names spelled from the profile, patience for a pause.",
        ]),
    ]
    gap = 0.22
    w = (WIDTH - 2 * gap) / 3
    for i, (question, points) in enumerate(cards):
        x = LEFT + i * (w + gap)
        rect(slide, x, 1.88, w, 3.45, CARD, radius=0.12, name=f"Card {i + 1}")
        text(slide, x + 0.2, 2.05, w - 0.4, 3.15,
             [P(R(question, 13, "semibold"), after=10, line=1.05)]
             + [P(R(point, 11), bullet=True, after=8, line=1.05) for point in points],
             name=f"Card {i + 1} text")


def features(slide):
    items = [
        ("MdRecordVoiceOver", "Talks in Japanese",
         "Gemini Live, voice to voice. Speaks first, asks one thing, waits 5 s through a pause."),
        ("MdHistory", "Remembers last time", "Opens by recalling the last conversation, or with a question the family sent."),
        ("MdMenuBook", "A life story, as a book",
         "Chapters by time of life. Each story retold, its quotes a press away, each with its clip."),
        ("MdFactCheck", "Never adds a detail",
         "Every fact retold has a quote; each quote found word for word, then heard in its clip."),
        ("MdTimeline", "Plans from the gaps", "Finds the times of life not yet told, and plans the next three questions."),
        ("MdVerifiedUser", "Asked at the close",
         "Shared only on a yes in their own words. No reply: asked again. Any story can be taken back."),
        ("MdForum", "The family asks back",
         "A relative says who they are and asks; it says who asked. Refused topics are never raised."),
        ("MdTranslate", "Japanese, with English",
         "The family page in one language at a time, with a switch. Quotes keep the words as said."),
    ]
    gap = 0.2
    w = (WIDTH - 3 * gap) / 4
    h = 1.82
    for i, (glyph, title, body) in enumerate(items):
        x = LEFT + (i % 4) * (w + gap)
        y = TOP + (i // 4) * (h + gap)
        rect(slide, x, y, w, h, CARD, radius=0.12, name=f"Feature {i + 1}")
        badge(slide, glyph, x + 0.18, y + 0.18)
        text(slide, x + 0.18, y + 0.78, w - 0.32, 0.95,
             [P(R(title, 12, "semibold"), after=3), P(R(body, 10, "regular", INK), line=1.05)],
             name=f"Feature {i + 1} text")


def process_flow(slide):
    """The five steps the start page's Under the hood shows, and the family's loop around them."""
    w, h = 2.6, 0.9
    xs = [LEFT, LEFT + w + 0.71, RIGHT - w]
    ys = [1.45, 2.78, 4.11]
    steps = {
        (0, 0): ("1", "Talk", "Live, by voice. Speaks first, asks one thing, waits through a pause."),
        (0, 1): ("2", "Write", "The talk becomes stories, each resting on their own sentences."),
        (0, 2): ("3", "Check", "Each quote found word for word, then heard in the recording."),
        (1, 2): ("4", "Share", "At the close: may the family read today's? Only a yes goes on."),
        (1, 1): ("5", "Plan", "Finds what is still untold; plans the next three questions."),
    }
    for (r, c), (n, title, body) in steps.items():
        rect(slide, xs[c], ys[r], w, h, CARD, radius=0.1, name=f"Step {n}")
        number(slide, n, xs[c] + 0.14, ys[r] + 0.14)
        text(slide, xs[c] + 0.56, ys[r] + 0.12, w - 0.7, h - 0.2,
             [P(R(title, 11.5, "semibold"), after=2), P(R(body, 9.5, "regular", MUTED), line=1.0)],
             name=f"Step {n} text")
    nodes = {
        (1, 0, CARD, "MdHistory", "Next conversation", "Opens with last time, or with the family's question."),
        (2, 2, ACCENT_TINT, "MdFamilyRestroom", "Family reads", "A book by time of life; each quote with its clip."),
        (2, 1, ACCENT_TINT, "MdForum", "Family asks back",
         "A relative says who they are, then asks: grandchild Yuka."),
        (2, 0, ACCENT_TINT, "MdVerifiedUser", "Screened", "A question touching a refused topic is quietly held back."),
    }
    for r, c, fill, glyph, title, body in sorted(nodes):
        rect(slide, xs[c], ys[r], w, h, fill, radius=0.1, name=title)
        badge(slide, glyph, xs[c] + 0.12, ys[r] + 0.12, d=0.34)
        text(slide, xs[c] + 0.56, ys[r] + 0.12, w - 0.7, h - 0.2,
             [P(R(title, 11.5, "semibold"), after=2), P(R(body, 9.5, "regular", MUTED), line=1.0)],
             name=f"{title} text")

    mid = [y + h / 2 for y in ys]
    centre = [x + w / 2 for x in xs]
    pad = 0.06
    arrow(slide, xs[0] + w + pad, mid[0], xs[1] - pad, mid[0])
    arrow(slide, xs[1] + w + pad, mid[0], xs[2] - pad, mid[0])
    arrow(slide, centre[2], ys[0] + h + pad, centre[2], ys[1] - pad)
    arrow(slide, xs[2] - pad, mid[1], xs[1] + w + pad, mid[1])
    arrow(slide, xs[1] - pad, mid[1], xs[0] + w + pad, mid[1])
    arrow(slide, centre[0], ys[1] - pad, centre[0], ys[0] + h + pad)
    arrow(slide, centre[2], ys[1] + h + pad, centre[2], ys[2] - pad, color=ACCENT)
    arrow(slide, xs[2] - pad, mid[2], xs[1] + w + pad, mid[2], color=ACCENT)
    arrow(slide, xs[1] - pad, mid[2], xs[0] + w + pad, mid[2], color=ACCENT)
    arrow(slide, centre[0], ys[2] - pad, centre[0], ys[1] + h + pad, color=ACCENT)
    text(slide, centre[2] - 1.62, ys[1] + h + 0.08, 1.5, 0.3,
         [P(R("only what was shared", 9, "medium", ACCENT_TEXT), align="r")], name="Shared label")
    text(slide, centre[0] + 0.12, ys[1] + h + 0.08, 1.5, 0.3,
         [P(R("two at most", 9, "medium", ACCENT_TEXT))], name="Limit label")
    text(slide, centre[0] + 0.12, ys[0] + h + 0.08, 1.5, 0.3,
         [P(R("next time", 9, "medium", MUTED))], name="Loop label")

    # Legend, beside the heading.
    rect(slide, 6.55, 0.93, 0.18, 0.18, CARD, radius=0.03, line=LINE, name="Legend swatch")
    text(slide, 6.8, 0.89, 1.5, 0.3, [P(R("Every conversation", 9.5, "medium", MUTED))], name="Legend 1")
    rect(slide, 8.35, 0.93, 0.18, 0.18, ACCENT_TINT, radius=0.03, line=ACCENT, name="Legend swatch 2")
    text(slide, 8.6, 0.89, 1.0, 0.3, [P(R("The family", 9.5, "medium", MUTED))], name="Legend 2")


def wireframes(slide):
    """The two screens as built: the senior's while talking, and the family page on a phone."""

    def callout(n, x, y):
        number(slide, n, x, y, d=0.24, fill=ACCENT, color=DARK)

    # Tablet, landscape: the senior's screen mid-conversation. Callouts in a column at its right.
    tx, ty, tw, th = LEFT, TOP, 4.75, 2.75
    rect(slide, tx, ty, tw, th, DARK, radius=0.22, name="Tablet frame")
    sx, sy, sw, sh = tx + 0.18, ty + 0.18, tw - 0.36, th - 0.36
    rect(slide, sx, sy, sw, sh, APP_BG, radius=0.06, name="Tablet screen")
    content = sw - 0.45
    cx = sx + content / 2
    text(slide, sx + 0.15, sy + 0.1, 2.0, 0.25, [P(R("小川節子さん", 9, "medium", ON_DARK_MUTED))])
    text(slide, sx + 0.25, sy + 0.32, content - 0.4, 0.78,
         [P(R("郡上おどり、いいですねえ。どなたと踊られたんですか。", 15, "medium", ON_DARK),
            align="ctr", line=1.1)],
         anchor="ctr", name="Interviewer's words")
    bars = [0.1, 0.22, 0.34, 0.2, 0.28, 0.14, 0.08]
    bx = cx - (len(bars) * 0.13 - 0.06) / 2
    for i, bh in enumerate(bars):
        rect(slide, bx + i * 0.13, sy + 1.36 - bh / 2, 0.07, bh, ACCENT, radius=0.035, name=f"Level {i}")
    text(slide, sx, sy + 1.56, content, 0.25,
         [P(R("お話ししています", 10, "medium", ON_DARK_MUTED), align="ctr")], name="Phase")
    stop = rect(slide, cx - 0.75, sy + 1.88, 1.5, 0.36, "A8503F", radius=0.18, name="Finish button")
    label_in(stop, [P(R("おわる", 13, "semibold", WHITE), align="ctr")])
    column = sx + sw - 0.36
    for n, centre in enumerate([sy + 0.7, sy + 1.45, sy + 2.06], 1):
        callout(n, column, centre - 0.12)

    text(slide, LEFT, ty + th + 0.12, tw, 0.3,
         [P(R("The senior's screen: one button to start, nothing to learn", 11, "semibold"))], name="Tablet caption")
    # Numbered as the screen's markers are; the close comes after, on a screen of its own.
    notes = ["Its words, large, as it says them", "Listening or speaking, calmly", "One button to finish",
             "Then, at the close: share today's story, or keep it"]
    for i, note in enumerate(notes):
        nx, ny = LEFT + (i % 2) * 2.4, ty + th + 0.5 + (i // 2) * 0.32
        if i < 3:
            number(slide, i + 1, nx, ny + 0.02, d=0.22, fill=ACCENT, color=DARK)
        else:
            text(slide, nx, ny - 0.02, 0.3, 0.28, [P(R("→", 11, "semibold", ACCENT_TEXT))], name="Tablet then")
        text(slide, nx + 0.3, ny, 2.05, 0.28, [P(R(note, 9.5, "regular", MUTED))], name=f"Tablet note {i + 1}")

    # Phone, portrait: the family page as a book. Callouts just outside its right edge.
    px, py, pw, ph = 5.4, TOP, 1.75, 3.6
    rect(slide, px, py, pw, ph, DARK, radius=0.22, name="Phone frame")
    qx, qy, qw, qh = px + 0.1, py + 0.18, pw - 0.2, ph - 0.36
    rect(slide, qx, qy, qw, qh, APP_BG, radius=0.08, name="Phone screen")
    text(slide, qx + 0.12, qy + 0.1, qw - 0.2, 0.62, [
        P(R("A LIFE STORY", 5.5, "semibold", ACCENT)),
        P(R("Ogawa Setsuko", 11, "semibold", ON_DARK)),
        P(R("小川節子 · Born in 1944, Gifu", 6, "regular", ON_DARK_MUTED)),
    ])
    for i, (label, x, wide) in enumerate([("A life story  3", qx + 0.12, 0.66), ("Questions  2", qx + 0.82, 0.6)]):
        pill = rect(slide, x, qy + 0.76, wide, 0.18, APP_PANEL if i else ACCENT, radius=0.09, name=label)
        label_in(pill, [P(R(label, 5.5, "semibold", ON_DARK if i else DARK), align="ctr")], inset=0.0)
    text(slide, qx + 0.12, qy + 1.04, qw - 0.2, 0.18, [P(R("CHAPTER 2 · MARRIAGE", 5.5, "semibold", ACCENT))])
    text(slide, qx + 0.12, qy + 1.2, qw - 0.2, 0.3,
         [P(R("Gujo Odori and Meeting Kazuo", 7.5, "semibold", ON_DARK), line=1.0)])
    text(slide, qx + 0.12, qy + 1.52, qw - 0.22, 0.4,
         [P(R("In the summer of Showa 40, when Setsuko was twenty-one, she met her husband Kazuo at the "
              "Gujo Odori…", 6, "regular", ON_DARK), line=1.05)])
    quotes = rect(slide, qx + 0.12, qy + 1.95, 0.78, 0.17, APP_BG, radius=0.08, line="5A4A3C", name="Quotes button")
    label_in(quotes, [P(R("All 2 quotes", 5.5, "medium", ACCENT), align="ctr")], inset=0.0)
    rect(slide, qx + 0.12, qy + 2.2, qw - 0.24, 0.8, APP_PANEL, radius=0.05, name="Quote")
    text(slide, qx + 0.2, qy + 2.24, qw - 0.4, 0.5, [
        P(R("「主人の和夫さんとはね、郡上おどりで知り合ったんですよ。…」", 6, "regular", ON_DARK), line=1.0, after=2),
        P(R("“I met my husband Kazuo…”", 5.5, "regular", ON_DARK_MUTED), line=1.0),
    ])
    circle(slide, qx + 0.2, qy + 2.76, 0.18, ACCENT, name="Play")
    icon(slide, "MdPlayArrow", qx + 0.222, qy + 2.782, 0.135, color=WHITE)
    rect(slide, qx + 0.44, qy + 2.845, qw - 0.64, 0.015, "5A4A3C", radius=0.0, name="Progress")
    centres = [qy + 0.38, qy + 0.85, qy + 2.6]
    for n, centre in enumerate(centres, 1):
        callout(n, px + pw + 0.06, centre - 0.12)

    family_notes = [
        "The life story as a book, by time of life",
        "The family's questions, one tab away, each with the story it led to",
        "Each retelling with its quotes, in the words as said, and the voice that said them",
    ]
    nx = px + pw + 0.4
    text(slide, nx, TOP, RIGHT - nx, 0.3, [P(R("The family page", 11, "semibold"))], name="Phone caption")
    for i, (note, centre) in enumerate(zip(family_notes, centres)):
        text(slide, nx, centre - 0.08, RIGHT - nx, 0.6, [P(R(note, 9.5, "regular", MUTED), line=1.0)],
             name=f"Family note {i + 1}")
    text(slide, nx, py + ph - 0.28, RIGHT - nx, 0.5,
         [P(R("Both are live at kataribe.kaseovo.com, with this family page and a conversation to try "
              "in English.", 9, "regular", MUTED), line=1.05)], name="Status note")


def architecture(slide):
    def node(x, y, w, h, glyph, title, lines, fill=CARD, name=None):
        rect(slide, x, y, w, h, fill, radius=0.1, name=name or title)
        badge(slide, glyph, x + 0.12, y + 0.12, d=0.36)
        text(slide, x + 0.56, y + 0.12, w - 0.66, h - 0.18,
             [P(R(title, 11, "semibold"), after=2)] + [P(R(line, 9, "regular", MUTED), line=1.0, after=2)
                                                      for line in lines],
             name=f"{title} text")

    cl_w = 1.95
    node(LEFT, TOP, cl_w, 2.35, "MdTabletAndroid", "Senior's screen",
         ["React web app", "Mic in, voice out at 24 kHz",
          "Records both voices on one timeline",
          "Holds only a single-use token, with the prompt sealed in"])
    node(LEFT, 4.3, cl_w, 1.05, "MdPhoneIphone", "Family page", ["Shared stories only, in Japanese or English"])

    run_x, run_w = 3.05, 3.5
    rect(slide, run_x, 2.05, run_w, 3.3, CARD, radius=0.12, name="Cloud Run")
    icon(slide, "MdCloud", run_x + 0.15, 2.12, 0.32)
    text(slide, run_x + 0.55, 2.1, run_w - 0.65, 0.42,
         [P(R("Cloud Run, Tokyo", 11, "semibold")),
          P(R("kataribe.kaseovo.com · FastAPI and the web app", 9, "regular", MUTED))],
         anchor="ctr", name="Cloud Run title")
    steps = [
        ("1", "Session start", "The story so far and family questions, sealed into a single-use Live token"),
        ("3", "Recording upload", "Stereo WAV in parts, joined in Cloud Storage"),
        ("4", "After each conversation", "Write → check → screen → plan"),
        ("5", "Family API", "Shared stories only; one clip per quote; English beside it"),
    ]
    sub_ys = [2.6, 3.28, 3.96, 4.64]
    for (n, title, body), y in zip(steps, sub_ys):
        rect(slide, run_x + 0.15, y, run_w - 0.3, 0.6, WHITE, radius=0.06, name=title)
        number(slide, n, run_x + 0.25, y + 0.15, d=0.3)
        text(slide, run_x + 0.65, y + 0.06, run_w - 0.85, 0.5,
             [P(R(title, 10, "semibold"), after=1), P(R(body, 8.5, "regular", MUTED), line=1.0)],
             anchor="ctr", name=f"{title} text")

    rx, rw = 7.2, RIGHT - 7.2
    node(rx, TOP, rw, 0.85, "MdGraphicEq", "Gemini Live API", ["gemini-3.8-live: native audio, built-in transcription"])
    services = [
        (sub_ys[0], "MdStorage", "Firestore", "People, sessions, stories, plans"),
        (sub_ys[1], "MdFolder", "Cloud Storage", "Recordings; a clip per quote"),
        (sub_ys[2], "MdAutoAwesome", "Gemini 3.8 Flash", "Writes, listens again, translates, plans"),
    ]
    for y, glyph, title, body in services:
        node(rx, y, rw, 0.6, glyph, title, [body])

    pad = 0.05
    # 2: the conversation, voice straight between the screen and Gemini, not
    # through Cloud Run: a step like the others, on the voice's own path.
    voice_y = 1.75
    arrow(slide, LEFT + cl_w + pad, voice_y, rx - pad, voice_y, color=ACCENT, width=2, both=True)
    rect(slide, run_x + 0.15, voice_y - 0.27, run_w - 0.3, 0.54, WHITE, radius=0.06, line=ACCENT, name="The conversation")
    number(slide, 2, run_x + 0.25, voice_y - 0.15, d=0.3, fill=ACCENT, color=DARK)
    text(slide, run_x + 0.65, voice_y - 0.24, run_w - 0.85, 0.48,
         [P(R("The conversation", 10, "semibold"), after=1),
          P(R("Voice both ways, screen to Gemini, on the token from 1", 8.5, "regular", MUTED), line=1.0)],
         anchor="ctr", name="The conversation text")
    for y in (sub_ys[0] + 0.3, sub_ys[1] + 0.3):
        arrow(slide, LEFT + cl_w + pad, y, run_x - pad, y, both=y == sub_ys[0] + 0.3)
    arrow(slide, LEFT + cl_w + pad, sub_ys[3] + 0.3, run_x - pad, sub_ys[3] + 0.3, both=True)
    for y in sub_ys[:3]:
        arrow(slide, run_x + run_w + pad, y + 0.3, rx - pad, y + 0.3, both=True)


def technologies(slide):
    groups = [
        ("MdAutoAwesome", "Google AI", [
            "Gemini 3.8 Live: real-time voice, native audio, built-in transcription",
            "Gemini 3.8 Flash: writes the stories, listens to each quote's clip, translates, plans",
            "Google Gen AI SDK, Python and JavaScript; single-use tokens",
        ]),
        ("MdCloud", "Google Cloud", [
            "Cloud Run in Tokyo, at kataribe.kaseovo.com",
            "Firestore and Cloud Storage; Secret Manager for the keys",
            "An uptime check, Error Reporting and a budget alert",
        ]),
        ("MdTabletAndroid", "Application", [
            "React 19, Vite, TypeScript",
            "Web Audio worklet: capture, and both voices on one recording",
            "FastAPI on Python 3.13, managed with uv",
        ]),
        ("MdFactCheck", "Quality", [
            "pytest and Playwright against the live API, with synthesized Japanese speech",
            "AI-vs-AI sparring: three personas, real audio both ways",
            "make smoke: a judge's whole path, on the live service",
        ]),
    ]
    gap = 0.2
    w = (WIDTH - 3 * gap) / 4
    for i, (glyph, title, points) in enumerate(groups):
        x = LEFT + i * (w + gap)
        rect(slide, x, TOP, w, 3.1, CARD, radius=0.12, name=title)
        badge(slide, glyph, x + 0.18, TOP + 0.18)
        text(slide, x + 0.18, TOP + 0.78, w - 0.32, 2.25,
             [P(R(title, 12.5, "semibold"), after=8)]
             + [P(R(point, 10.5), bullet=True, after=7, line=1.05) for point in points],
             name=f"{title} text")
    rect(slide, LEFT, 4.72, WIDTH, 0.62, DARK, radius=0.12, name="Principle")
    text(slide, LEFT + 0.25, 4.76, WIDTH - 0.5, 0.54, [
        P(R("One rule across the stack  ", 11, "semibold", ACCENT),
          R("The model never speaks for the senior. Gemini writes the stories; plain code checks every quote "
            "and every yes against the transcript, and Gemini hears each one again in the recording.", 10,
            "regular", ON_DARK), line=1.05),
    ], anchor="ctr", name="Principle text")


def cost(slide):
    rect(slide, LEFT, TOP, 3.0, 3.88, DARK, radius=0.12, name="Cost summary")
    text(slide, LEFT + 0.28, TOP + 0.25, 2.5, 3.5, [
        P(R("≈ $0.13", 34, "semibold", ACCENT), after=2),
        P(R("per 10-minute conversation, in AI and cloud, at list prices", 11.5, "medium", ON_DARK), after=16, line=1.05),
        P(R("≈ ¥6,900", 26, "semibold", ON_DARK), after=2),
        P(R("for a year of daily conversations, per person", 11.5, "medium", ON_DARK), after=16, line=1.05),
        P(R("About 3% of one ¥220,000 self-history written by hand.", 10.5, "regular", ON_DARK_MUTED), line=1.05),
    ], name="Cost summary text")

    rows = [
        ["Item", "Basis", "Per session"],
        ["Gemini Live, audio in", "$0.005 a minute, for all 10 minutes", "$0.050"],
        ["Gemini Live, audio out", "$0.018 a minute, about 2.5 minutes spoken", "$0.045"],
        ["Gemini Flash, writing", "≈ 15k tokens in, 3k out: write, screen, plan", "$0.023"],
        ["Gemini Flash, listening", "≈ 3 minutes of quoted clips, heard again", "≈ $0.006"],
        ["Recording storage", "58 MB of WAV, at ≈ $0.023 per GB-month", "< $0.002"],
        ["Cloud Run, Firestore", "Scale to zero; free tiers at this scale", "≈ $0"],
        ["Total", "", "≈ $0.13", "*"],
    ]
    table(slide, 3.65, TOP, [1.85, 3.0, 1.11], rows, row_h=0.38)
    text(slide, 3.65, TOP + 3.16, RIGHT - 3.65, 0.75, [
        P(R("Audio is billed for the whole time a session is open, silence included. Rates are Google's list "
            "prices as of 23 Sep 2026, kept in backend/costs.py; token counts are estimates, the listening at 32 "
            "audio tokens a second. Yen at ¥150 to the dollar. Judging: at most 30 visitor conversations a day, "
            "about $0.10 each.", 9, "regular", MUTED), line=1.05),
    ], name="Cost notes")


def snapshots(slide):
    shots = [
        ("start.png", "The start page",
         "What it is, who it's for, a family page from the demo, then a judge's own turn."),
        ("story.png", "A story, and what it rests on",
         "A retelling to read; every quote a press away, in Japanese with the English beneath."),
        ("questions.png", "Questions from the family",
         "Who asked, and the story each question led to. A waiting one says so."),
    ]
    gap = 0.31
    w = (WIDTH - 2 * gap) / 3
    h = w * 480 / 430
    for i, (file, title, body) in enumerate(shots):
        x = LEFT + i * (w + gap)
        picture(slide, SHOTS / file, x, TOP - 0.05, w, h, radius=0.1)
        text(slide, x, TOP + h + 0.04, w, 0.7,
             [P(R(title, 11, "semibold"), after=1), P(R(body, 9, "regular", MUTED), line=1.0)],
             name=f"{title} caption")
    text(slide, 5.2, 0.86, RIGHT - 5.2, 0.42,
         [P(R("Live at kataribe.kaseovo.com, on a phone: the demo's family page, written from the video's conversations.",
              9, "regular", MUTED), align="r", line=1.0)],
         name="Data note")


def benchmarks(slide):
    tiles = [
        ("36 / 36", "Consent, twelve cases: a yes shares; a no or a hesitation keeps it private; no reply is asked again"),
        ("12 / 12", "The recording check: the same words confirmed, 18 for 十八 accepted, a changed place caught"),
        ("3 / 3", "A judge's whole path on the live service, from the start page to their own family page"),
        ("6 min", "A browser session with no drop, 34 transcript entries"),
    ]
    gap = 0.2
    w = (WIDTH - 3 * gap) / 4
    for i, (big, body) in enumerate(tiles):
        x = LEFT + i * (w + gap)
        rect(slide, x, TOP, w, 1.45, DARK, radius=0.12, name=f"Result {big}")
        text(slide, x + 0.2, TOP + 0.14, w - 0.35, 1.25,
             [P(R(big, 24, "semibold", ACCENT), after=4), P(R(body, 9.5, "regular", ON_DARK), line=1.05)],
             name=f"Result {big} text")

    y = TOP + 1.65
    text(slide, LEFT, y, 4.3, 0.3, [P(R("How long it waits before replying", 11.5, "semibold"))], name="Patience title")
    rows = [
        ["Patience setting", "Pause in speech", "Result"],
        ["3,500 ms", "1.5 s", "Waited"],
        ["3,500 ms", "2.5 s", "Cut in"],
        ["6,000 ms", "2.5 s", "Waited"],
        ["6,000 ms", "4.0 s", "Waited"],
    ]
    table(slide, LEFT, y + 0.35, [1.55, 1.45, 1.3], rows, row_h=0.3, size=10)
    text(slide, LEFT, y + 1.92, 4.3, 0.3,
         [P(R("Japanese: 5,000 ms, measured. English visitors: 4 s. The model's own: 800 ms.", 9.5, "medium",
              ACCENT_TEXT))],
         name="Patience note")

    method = [
        "Real models, three runs a case: speech streamed at the Live API in real time, replies graded by rules "
        "and by a model judging them against the prompt.",
        "AI-vs-AI sparring: three personas press a long pause, a refused topic, a change of subject.",
        "The recording check caught a real slip on its own: 私は transcribed as 私を. The quote was held back.",
        "Not yet measured: elderly voices and dialect. A native speaker's review of the Japanese is next.",
    ]
    text(slide, 5.0, y, RIGHT - 5.0, 2.1,
         [P(R("How it is measured", 11.5, "semibold"), after=6)]
         + [P(R(point, 10), bullet=True, after=5, line=1.05) for point in method],
         name="Method")


def future(slide):
    columns = [
        ("MdSensors", "Next, on the same screen", [
            "It notices the senior sitting nearby, so there is no button at all",
            "Names it is unsure of kept in kana; the family fixes the spelling in one tap",
            "A message to the family when a new story lands",
        ]),
        ("MdPhoneInTalk", "New channels, the same brain", [
            "Phone and landline",
            "Smart speakers",
            "Inside the companion robots Japanese families already buy",
        ]),
        ("MdGroups", "Beyond the family", [
            "The teller chooses who else may read: a friend, a care worker, a local archive",
            "Testimony, the oldest sense of 語り部: the war, the bombings, the disasters, told by those who lived them",
            "Municipalities as payers: Yokosuka City has trialled conversation AI for seniors since 2025",
        ]),
    ]
    gap = 0.22
    w = (WIDTH - 2 * gap) / 3
    for i, (glyph, title, points) in enumerate(columns):
        x = LEFT + i * (w + gap)
        rect(slide, x, TOP, w, 2.75, CARD, radius=0.12, name=title)
        badge(slide, glyph, x + 0.18, TOP + 0.18)
        text(slide, x + 0.18, TOP + 0.78, w - 0.34, 1.9,
             [P(R(title, 12, "semibold"), after=7)]
             + [P(R(point, 10.5), bullet=True, after=5, line=1.05) for point in points],
             name=f"{title} text")
    rect(slide, LEFT, 4.4, WIDTH, 0.93, DARK, radius=0.12, name="Scale")
    text(slide, LEFT + 0.25, 4.5, WIDTH - 0.5, 0.75, [
        P(R("How it scales  ", 11.5, "semibold", ACCENT),
          R("Every channel plugs into the same backend. Cloud Run scales to zero and out. Gemini moves to "
            "Vertex AI in a Japanese region, behind one adapter, when data residency matters.", 10.5,
            "regular", ON_DARK), line=1.05),
    ], anchor="ctr", name="Scale text")


def links(slide):
    box = heading_box(slide)
    values = {
        "GitHub Public Repository": REPO_URL,
        "Demo Video Link (3 Minutes)": VIDEO_URL,
        "Final Product Link": PRODUCT_URL,
    }
    for paragraph in box.text_frame.paragraphs:
        if not paragraph.runs:
            continue
        label = paragraph.runs[0].text
        if label in values:
            value = values[label]
            _write_run(paragraph, R(":  ", 14, "medium", INK))
            if value:
                _write_run(paragraph, R(value.removeprefix("https://"), 14, "medium", ACCENT_TEXT, link=value))
            else:
                _write_run(paragraph, R("TBD", 14, "semibold", "C5221F"))
            end = paragraph._p.find(qn("a:endParaRPr"))
            if end is not None:
                paragraph._p.remove(end)
                paragraph._p.append(end)
    box.height = emu(1.6)

    import qrcode

    blocks = [
        (VIDEO_URL, "qr-video.png", "Watch the demo",
         "3:48, in English: two conversations, the family page they made, and how to try it yourself."),
        (PRODUCT_URL, "qr-product.png", "Try it on your own life",
         "A family page from the demo, then a conversation of your own, in English. About three minutes."),
        (REPO_URL, "qr-repo.png", "The code, the tests, and the plan behind them",
         "Open source, under the MIT licence. Built in the open from the first commit on 22 Sep 2026."),
    ]
    for i, (url, file, title, body) in enumerate(blocks):
        code = qrcode.QRCode(border=0, box_size=20)
        code.add_data(url)
        path = HERE / "assets" / file
        code.make_image(fill_color="#1F1F1F", back_color="white").save(path)
        x = LEFT + 0.1 + i * 3.1
        picture(slide, path, x, 2.85, 1.2, 1.2)
        text(slide, x + 1.35, 2.7, 1.6, 1.5, [
            P(R(title, 12, "semibold"), after=4),
            P(R(body, 10, "regular", MUTED), line=1.05),
        ], anchor="ctr", name=f"{title} note")


def business_case(slide, heading_source):
    heading = copy.deepcopy(heading_source._element)
    slide.shapes._spTree.append(heading)
    run = heading.find(".//" + qn("a:r"))
    run.find(qn("a:t")).text = "Community benefit and business case"

    columns = [
        ("MdElderly", "For them", [
            "A few minutes a day that are theirs, about their own life",
            "Remembering out loud, with someone who waits for a name or a year",
            "Theirs to give: nothing reaches the family until they say so",
        ]),
        ("MdFamilyRestroom", "For their family", [
            "The stories they never asked for: not for lack of love, for lack of a moment",
            "In their own words, with their voice behind each quote",
            "Something to talk about, and a question to send",
        ]),
        ("MdGroups", "For their town, next", [
            "Whoever they choose: a friend, a care worker, a local archive",
            "Many of the 6.72 million who live alone have nobody nearby; no family is needed",
            "What a town lived through, in the voices of those who lived it",
        ]),
    ]
    gap = 0.22
    w = (WIDTH - 2 * gap) / 3
    for i, (glyph, title, points) in enumerate(columns):
        x = LEFT + i * (w + gap)
        rect(slide, x, TOP, w, 2.55, CARD, radius=0.12, name=title)
        badge(slide, glyph, x + 0.18, TOP + 0.18)
        text(slide, x + 0.18, TOP + 0.78, w - 0.34, 1.7,
             [P(R(title, 12, "semibold"), after=7)]
             + [P(R(point, 10.5), bullet=True, after=5, line=1.05) for point in points],
             name=f"{title} text")

    payers = [
        ("Families", "A subscription, against a ¥220,000 one-off book"),
        ("Municipalities and care providers", "Preventive care and reminiscence programmes"),
        ("Cost to serve", "≈ ¥570 a month for a daily conversation"),
    ]
    y = 4.2
    rect(slide, LEFT, y, WIDTH, 1.13, DARK, radius=0.12, name="Who pays")
    text(slide, LEFT + 0.25, y + 0.12, 1.3, 0.9, [P(R("Who pays", 12, "semibold", ACCENT))], anchor="ctr")
    pw = (WIDTH - 1.75) / 3
    for i, (who, how) in enumerate(payers):
        text(slide, LEFT + 1.55 + i * pw, y + 0.12, pw - 0.2, 0.9,
             [P(R(who, 10.5, "semibold", ON_DARK), after=2), P(R(how, 10, "regular", ON_DARK_MUTED), line=1.05)],
             anchor="ctr", name=f"Payer {i + 1}")


def main() -> None:
    prs = Presentation(HERE / "template.pptx")
    remove_slide(prs, 1)  # "Please download this template…": an instruction, not a slide
    s = list(prs.slides)
    team_details(s[0])
    brief(s[1])
    opportunities(s[2])
    features(s[3])
    process_flow(s[4])
    wireframes(s[5])
    architecture(s[6])
    technologies(s[7])
    cost(s[8])
    snapshots(s[9])
    benchmarks(s[10])
    future(s[11])
    links(s[12])
    business_case(s[13], heading_box(s[1]))
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
